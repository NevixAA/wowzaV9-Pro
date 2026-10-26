"""Deep Wowza analytics snapshot -> Excel workbook + Markdown report + CSVs (owner brief 2026-10-18).

    python -m src.analytics.deep_wowza_snapshot [--out DIR]

READ ONLY. Reads v9's ledgers and policy (V9_LOCAL), Pro's research artifacts and V11's public outputs;
writes only output/deep_analytics/<date>/. It changes no selection logic, threshold, model, staking,
league eligibility or production behaviour.

Main window 2026-08-26 -> latest settled date; every sheet states its window. Bets are canonical
FIRST-SIGHT records (src/analytics/core.py). Primary result convention: 1 unit flat stake per bet.
"""
from __future__ import annotations

import argparse
import io
import json
from pathlib import Path

import numpy as np
import pandas as pd

from config import pro_config as cfg
from src.analytics import core as C

MAIN_START = "2026-08-26"
RECENT_FROM = "2026-09-24"
V11_RAW = "https://raw.githubusercontent.com/NevixAA/wowza_v11/main/output/"
DISCLAIMER = ("These are run-rate scenarios, not forecasts. They assume observed volume and/or ROI persists and "
              "do not model future fixture calendar, market adaptation, injuries, model drift, variance, or "
              "regression to the mean.")
PCT_COLS = ("hit_rate", "roi", "avg_implied_probability", "break_even_hit_rate", "avg_edge", "median_edge",
            "avg_model_probability", "median_model_probability", "avg_clv", "median_clv", "pct_positive_clv",
            "recent_ROI", "last30_ROI", "last14_ROI", "last7_ROI", "roi_ci_lo_90", "roi_ci_hi_90", "roi_ci_lo_95",
            "roi_ci_hi_95", "hit_ci_lo_90", "hit_ci_hi_90", "current_ROI", "recent_ROI", "shrunk_ROI",
            "running_roi", "share_of_bets", "share_of_positive_profit", "share_of_pnl", "toward_rate", "toward_ci_lo",
            "toward_ci_hi", "mean_clv_pct", "median_clv_pct", "pct_positive_clv_v11", "model_predicted", "actual",
            "bookmaker_implied", "calibration_error")


# ── inputs ───────────────────────────────────────────────────────────────────────────────────
def _json(p: Path) -> dict:
    try:
        return json.loads(Path(p).read_text(encoding="utf-8"))
    except Exception:                                                # noqa: BLE001
        return {}


def _v11(name: str) -> pd.DataFrame:
    local = cfg.BASE_DIR.parent / "wowza_v11" / "output" / name
    try:
        if local.exists():
            return pd.read_csv(local)
        import requests
        r = requests.get(V11_RAW + name, timeout=60)
        return pd.read_csv(io.StringIO(r.text)) if r.status_code == 200 else pd.DataFrame()
    except Exception:                                                # noqa: BLE001
        return pd.DataFrame()


def load_inputs() -> dict:
    v9 = cfg.V9_LOCAL
    raw = C.load_raw(v9)
    fs, stats = C.first_sight(raw)
    s = C.settle(fs)
    policy = _json(v9 / "config" / "bankroll_policy.json")
    return {"raw": raw, "first": fs, "stats": stats, "s": s, "policy": policy,
            "end": s["date"].max(), "retrain": _json(v9 / "output" / "retrain_log.json"),
            "bankroll_rec": _json(cfg.OUTPUT_DIR / "bankroll" / "current_policy_recommendation.json"),
            "pure": _json(cfg.OUTPUT_DIR / "pure_prediction" / "verdicts.json"),
            "player_v3": _json(cfg.OUTPUT_DIR / "player_v3" / "player_v3_verdict.json"),
            "scout": _json(cfg.OUTPUT_DIR / "scout" / "weekly_promotion_review.json"),
            "scout_status": _json(cfg.OUTPUT_DIR / "scout" / "league_status.json"),
            "tip_ranker": _json(cfg.OUTPUT_DIR / "studies" / "tip_ranker.json"),
            "v11_league": _v11("v11_movement_by_league.csv"), "v11_summary": _v11("v11_movement_summary.csv"),
            "v11_placebo": _v11("v11_placebo_table.csv"), "v11_residual": _v11("v11_residual.csv"),
            "v11_time": _v11("v11_movement_by_time.csv"), "v11_health": _v11_json("v11_research_health.json")}


def _v11_json(name: str) -> dict:
    try:
        import requests
        r = requests.get(V11_RAW + name, timeout=60)
        return r.json() if r.status_code == 200 else {}
    except Exception:                                                # noqa: BLE001
        return {}


# ── tables ───────────────────────────────────────────────────────────────────────────────────
def build_tables(I: dict) -> dict[str, pd.DataFrame]:
    s, end, pol = I["s"], I["end"], I["policy"]
    U = C.universes(s, cfg.PAPER_LEAGUES, pol)
    W = C.windows(end)
    main = W[f"since_{MAIN_START}"]
    core, nonmls = C.in_window(U["CORE"], main), C.in_window(U["NON_MLS"], main)
    tracks = pol.get("tracks", {})
    def state(t):
        """Policy state of model|market. O/U 2.5 is per side in v9 since 2026-10-26: one state when both
        sides agree, else 'OVER:<s> / UNDER:<s>'."""
        if t in tracks:
            return tracks[t].get("state", "UNLISTED")
        o, u = tracks.get(f"{t}_OVER"), tracks.get(f"{t}_UNDER")
        if o is None and u is None:
            return "UNLISTED"
        so, su = (o or {}).get("state", "UNLISTED"), (u or {}).get("state", "UNLISTED")
        return so if so == su else f"OVER:{so} / UNDER:{su}"
    T = {}
    tab = lambda g, by, ci=True: C.table(g, by, end, RECENT_FROM, ci)  # noqa: E731
    win_txt = f"{MAIN_START} -> {end.date()}"

    # Overall: universes x windows
    rows = []
    for uname, ug in U.items():
        for wname, w in W.items():
            m = C.summarize(C.in_window(ug, w))
            rows.append({"universe": uname, "window": wname, "from": str(w[0].date()), "to": str(w[1].date()), **m})
    T["Overall"] = pd.DataFrame(rows)

    mm = []
    for uname, ug in (("CORE", core), ("NON_MLS_ALL_TIERS", nonmls)):
        t = tab(ug, ["model", "market"])
        t.insert(0, "universe", uname)
        t.insert(3, "policy_state", (t["model"] + "|" + t["market"]).map(state))
        for c, v in t.iterrows():
            p = C.projection(ug[(ug["model"] == v["model"]) & (ug["market"] == v["market"])], MAIN_START, RECENT_FROM, end)
            t.at[c, "season_tip_estimate"] = p["full_window_projected_tips"]
            t.at[c, "season_PnL_if_full_window_holds"] = p["full_window_projected_PnL"]
            t.at[c, "season_PnL_if_recent_holds"] = p["recent_projected_PnL"]
        mm.append(t)
    T["Market_Model"] = pd.concat(mm, ignore_index=True)

    tier = [tab(nonmls, ["tier"]).assign(view="tier"), tab(nonmls, ["model", "tier"]).assign(view="model x tier"),
            tab(nonmls, ["market", "tier"]).assign(view="market x tier")]
    T["Tier"] = pd.concat(tier, ignore_index=True)
    T["Model_Market_Tier"] = tab(nonmls, ["model", "market", "tier"])

    lg = tab(nonmls, ["league"])
    extra = []
    for _, r in lg.iterrows():
        g = nonmls[nonmls["league"] == r["league"]]
        bym = g.groupby("market")["pnl"].agg(["sum", "size"]); bym = bym[bym["size"] >= 5]
        bymo = g.groupby("model")["pnl"].agg(["sum", "size"])
        byt = g.groupby("tier")["pnl"].agg(["sum", "size"])
        cells = g.groupby(["model", "market", "tier"])["pnl"].sum()
        extra.append({"market_distribution": "; ".join(f"{k}:{int(v)}" for k, v in g["market"].value_counts().items()),
                      "model_distribution": "; ".join(f"{k}:{int(v)}" for k, v in g["model"].value_counts().items()),
                      "tier_distribution": "; ".join(f"{k}:{int(v)}" for k, v in g["tier"].value_counts().items()),
                      "best_market": (bym["sum"] / bym["size"]).idxmax() if len(bym) else "",
                      "worst_market": (bym["sum"] / bym["size"]).idxmin() if len(bym) else "",
                      "best_model": (bymo["sum"] / bymo["size"]).idxmax() if len(bymo) else "",
                      "best_tier": (byt["sum"] / byt["size"]).idxmax() if len(byt) else "",
                      "largest_profit_cell": " | ".join(cells.idxmax()) + f" ({cells.max():+.2f}u)" if len(cells) else "",
                      "largest_loss_cell": " | ".join(cells.idxmin()) + f" ({cells.min():+.2f}u)" if len(cells) else ""})
    T["League"] = pd.concat([lg, pd.DataFrame(extra)], axis=1)
    T["League_Market"] = tab(nonmls, ["league", "model", "market", "tier"])
    cells = tab(nonmls, ["league", "model", "market", "tier"])
    for c, v in cells.iterrows():
        g = nonmls[(nonmls["league"] == v["league"]) & (nonmls["model"] == v["model"]) &
                   (nonmls["market"] == v["market"]) & (nonmls["tier"] == v["tier"])]
        p = C.projection(g, MAIN_START, RECENT_FROM, end)
        cells.at[c, "season_equivalent_tips"] = p["full_window_projected_tips"]
        cells.at[c, "season_equivalent_pnl"] = p["full_window_projected_PnL"]
        cells.at[c, "recent_regime_season_tips"] = p["recent_projected_tips"]
        cells.at[c, "recent_regime_season_pnl"] = p["recent_projected_PnL"]
        cells.at[c, "policy_state"] = state(f"{v['model']}|{v['market']}")
    T["Detailed_Cells"] = cells

    nm = nonmls.assign(odds_band=C.band(nonmls["odds"], C.ODDS_BANDS), edge_band=C.band(nonmls["edge"], C.EDGE_BANDS))
    T["Odds_Bands"] = pd.concat([tab(nm, ["model", "market", "odds_band"]).assign(view="model x market"),
                                 tab(nm, ["model", "market", "tier", "odds_band"]).assign(view="model x market x tier"),
                                 tab(nm, ["odds_band"], ci=True).assign(view="all non-MLS")], ignore_index=True)
    eb = nm[nm["edge"].notna()]
    T["Edge_Bands"] = pd.concat([tab(eb, ["model", "market", "edge_band"]).assign(view="model x market"),
                                 tab(eb, ["edge_band"]).assign(view="all non-MLS")], ignore_index=True)
    rk = []
    for (mo, mk), g in eb.groupby(["model", "market"]):
        if len(g) >= 20:
            rk.append({"model": mo, "market": mk, "N": len(g),
                       "spearman_edge_vs_return": g["edge"].corr(g["pnl"], method="spearman"),
                       "spearman_edge_vs_win": g["edge"].corr(g["win"], method="spearman"),
                       "spearman_edge_vs_clv": g["edge"].corr(g["clv"], method="spearman") if g["clv"].notna().sum() >= 20 else np.nan})
    T["Edge_Rank"] = pd.DataFrame(rk)

    ou = nonmls[nonmls["market"] == "OU25"]
    T["OU25_Side"] = pd.concat([tab(ou, ["model", "side"]).assign(view="model x side"),
                                tab(ou, ["model", "side", "tier"]).assign(view="model x side x tier"),
                                tab(ou, ["league", "model", "side"]).assign(view="league x model x side")], ignore_index=True)

    tracks6 = {"core": core, "new_format": core[core.model == "new_format"], "standard": core[core.model == "standard"],
               "nf_BTTS": core[(core.model == "new_format") & (core.market == "BTTS")],
               "nf_OU25": core[(core.model == "new_format") & (core.market == "OU25")],
               "std_OU25": core[(core.model == "standard") & (core.market == "OU25")]}
    mon = []
    for name, g in [("overall_core", core), ("new_format", tracks6["new_format"]), ("standard", tracks6["standard"]),
                    ("OU25", core[core.market == "OU25"]), ("BTTS", core[core.market == "BTTS"]),
                    ("OVER15", core[core.market == "OVER15"]), ("SNIPER", core[core.tier == "SNIPER"]),
                    ("MARKSMAN", core[core.tier == "MARKSMAN"]), ("VALUABLE", C.in_window(U["VALUABLE"], main))]:
        t = tab(g, ["month"], ci=False).assign(track=name)
        mon.append(t)
    T["Monthly"] = pd.concat(mon, ignore_index=True)

    wk = tab(core, ["week"], ci=False)
    for lab, mask in (("new_format_ROI", core.model == "new_format"), ("standard_ROI", core.model == "standard"),
                      ("BTTS_ROI", core.market == "BTTS"), ("OU25_ROI", core.market == "OU25"),
                      ("SNIPER_ROI", core.tier == "SNIPER"), ("MARKSMAN_ROI", core.tier == "MARKSMAN")):
        wk[lab] = wk["week"].map(core[mask].groupby("week")["pnl"].mean())
    T["Weekly"] = wk

    dy = core.groupby("day").agg(N=("pnl", "size"), W=("win", "sum"), L=("loss", "sum"), PnL=("pnl", "sum")).reset_index()
    dy["ROI"] = dy["PnL"] / dy["N"]
    for lab, mask in (("new_format", core.model == "new_format"), ("standard", core.model == "standard")):
        a = core[mask].groupby("day")["pnl"].agg(["size", "sum"])
        dy[f"{lab}_N"], dy[f"{lab}_PnL"] = dy["day"].map(a["size"]).fillna(0), dy["day"].map(a["sum"]).fillna(0)
    dy["BTTS_PnL"] = dy["day"].map(core[core.market == "BTTS"].groupby("day")["pnl"].sum()).fillna(0)
    dy["OU_PnL"] = dy["day"].map(core[core.market == "OU25"].groupby("day")["pnl"].sum()).fillna(0)
    T["Daily"] = dy

    T["Rolling"] = pd.concat([C.rolling(g, n) for n, g in tracks6.items() if len(g)], ignore_index=True)
    T["Profit_Curves"] = pd.concat([C.curve(g, n) for n, g in tracks6.items() if len(g)], ignore_index=True)
    dd = [dict(track=n, **C.drawdown(g)) for n, g in list(tracks6.items()) +
          [("SNIPER", core[core.tier == "SNIPER"]), ("MARKSMAN", core[core.tier == "MARKSMAN"])]]
    T["Drawdown"] = pd.DataFrame(dd)

    # concentration
    def conc(by):
        t = nonmls.groupby(by).agg(N=("pnl", "size"), PnL=("pnl", "sum")).reset_index()
        pos = t["PnL"].clip(lower=0).sum()
        t["ROI"] = t["PnL"] / t["N"]
        t["share_of_bets"] = t["N"] / t["N"].sum()
        t["share_of_positive_profit"] = t["PnL"].clip(lower=0) / pos if pos else np.nan
        t["share_of_pnl"] = t["PnL"] / t["PnL"].sum() if t["PnL"].sum() else np.nan
        t["flag"] = np.where(t["share_of_positive_profit"] > 0.4, "CONCENTRATED (>40% of positive profit)", "")
        return t.assign(dimension=by).rename(columns={by: "value"})
    T["Concentration"] = pd.concat([conc("league"), conc("market"), conc("model"), conc("tier")], ignore_index=True)

    ex = []
    for name, g in (("core", core), ("nf_BTTS", tracks6["nf_BTTS"]), ("new_format", tracks6["new_format"]),
                    ("nf_BTTS_recent", tracks6["nf_BTTS"][tracks6["nf_BTTS"].date >= RECENT_FROM]),
                    ("new_format_recent", tracks6["new_format"][tracks6["new_format"].date >= RECENT_FROM])):
        a, b = C.summarize(g, ci=False), C.summarize(g[g.league != "Argentina Primera Division"], ci=False)
        top = g.groupby("league")["pnl"].sum()
        tl = top.idxmax() if len(top) else ""
        c = C.summarize(g[g.league != tl], ci=False) if tl else {}
        ex.append({"track": name, "N": a.get("N"), "ROI": a.get("roi"), "PnL": a.get("pnl_units"),
                   "exARG_N": b.get("N"), "exARG_ROI": b.get("roi"), "exARG_PnL": b.get("pnl_units"),
                   "top_profit_league": tl, "exTOP_N": c.get("N"), "exTOP_ROI": c.get("roi"), "exTOP_PnL": c.get("pnl_units")})
    T["Ex_Argentina_Top"] = pd.DataFrame(ex)

    # CLV
    clvd = []
    for by in ("model", "market", "tier", "league"):
        for k, g in nm.groupby(by):
            c = g["clv"].dropna()
            q = c.quantile([0.05, 0.25, 0.75, 0.95]) if len(c) else pd.Series([np.nan] * 4, index=[0.05, 0.25, 0.75, 0.95])
            clvd.append({"dimension": by, "value": k, "N": len(g), "clv_availability": len(c) / len(g),
                         "mean": c.mean(), "median": c.median(), "positive": (c > 0).mean() if len(c) else np.nan,
                         "p05": q[0.05], "p25": q[0.25], "p75": q[0.75], "p95": q[0.95],
                         "extreme_values_abs_gt_50pct": int((c.abs() > 0.5).sum())})
    T["CLV"] = pd.DataFrame(clvd)
    cv = cells[(cells["N"] >= 25) & cells["avg_clv"].notna()][["league", "model", "market", "tier", "N", "roi", "avg_clv"]]
    T["CLV_vs_ROI"] = cv.assign(corr_note=f"Spearman ROI~CLV over {len(cv)} cells (N>=25): "
                                f"{cv['roi'].corr(cv['avg_clv'], method='spearman'):.2f}" if len(cv) >= 5 else "too few cells")

    cal = nm[nm["model_p"].notna()].assign(prob_band=lambda x: C.band(x["model_p"], C.PROB_BANDS))
    cl = cal.groupby(["model", "market", "prob_band"]).agg(N=("win", "size"), model_predicted=("model_p", "mean"),
                                                            actual=("win", "mean"), bookmaker_implied=("implied", "mean"),
                                                            roi=("pnl", "mean")).reset_index()
    cl["calibration_error"] = cl["model_predicted"] - cl["actual"]
    T["Calibration"] = cl

    vol = nonmls.groupby("day").agg(all_tiers=("pnl", "size"))
    vol["core"] = core.groupby("day")["pnl"].size()
    for mk in ("OU25", "BTTS", "OVER15", "OVER35"):
        vol[mk] = nonmls[nonmls.market == mk].groupby("day")["pnl"].size()
    vol = vol.fillna(0)
    stats = vol.describe(percentiles=[0.5, 0.75, 0.9]).T[["mean", "50%", "75%", "90%", "max"]]
    T["Daily_Volume"] = pd.concat([stats.reset_index().rename(columns={"index": "series", "50%": "median", "75%": "p75", "90%": "p90"}),
                                   vol.sort_values("all_tiers", ascending=False).head(10).reset_index().assign(series="top-volume days")],
                                  ignore_index=True)

    proj = []
    for name, g in [("overall core", core), ("all non-MLS", nonmls), ("new_format", tracks6["new_format"]),
                    ("standard", tracks6["standard"]), ("BTTS new_format", tracks6["nf_BTTS"]),
                    ("OU25 new_format", tracks6["nf_OU25"]),
                    ("O1.5 new_format", core[(core.model == "new_format") & (core.market == "OVER15")]),
                    ("OU25 standard", tracks6["std_OU25"]), ("SNIPER", core[core.tier == "SNIPER"]),
                    ("MARKSMAN", core[core.tier == "MARKSMAN"]), ("VALUABLE", C.in_window(U["VALUABLE"], main))] + \
            [(f"league: {k}", g) for k, g in nonmls.groupby("league") if len(g) >= 50]:
        proj.append({"row": name, **C.projection(g, MAIN_START, RECENT_FROM, end)})
    T["Season_Projection"] = pd.DataFrame(proj)

    # policy
    pr, pb = [], []
    for t, v in tracks.items():
        pr.append({"track": t, "POLICY_STATE": v.get("state"), "BASE_FRACTION": v.get("base_fraction"),
                   "MIN_ODDS": v.get("min_odds"), "MAX_ODDS": v.get("max_odds"), "evidence_score": v.get("evidence_score")})
        g = core[core["track"] == t]
        if g.empty:
            continue
        lo, hi = float(v.get("min_odds") or 1.0), float(v.get("max_odds") or 99)
        ins = g[(g.odds >= lo) & (g.odds <= hi)]
        for lab, x in (("all signals", g), ("within policy price band", ins), ("excluded by price band", g.drop(ins.index))):
            m = C.summarize(x, ci=False)
            pb.append({"track": t, "state": v.get("state"), "subset": lab, "N": m.get("N"), "roi": m.get("roi"),
                       "pnl_units": m.get("pnl_units"), "avg_clv": m.get("avg_clv"),
                       "max_drawdown_units": C.drawdown(x).get("max_drawdown_units") if len(x) else np.nan})
    pinfo = [{"item": "policy_version", "value": pol.get("policy_version")}, {"item": "status", "value": pol.get("status")},
             {"item": "tier_multipliers", "value": json.dumps(pol.get("tier_multipliers"))},
             {"item": "portfolio caps", "value": json.dumps(pol.get("portfolio"))},
             {"item": "league_rating_sizing", "value": json.dumps(pol.get("league_rating_sizing"))}]
    rv = (I["bankroll_rec"] or {}).get("review", {})
    for k in ("RECOMMENDED_POLICY", "POLICY_CHANGE_RECOMMENDED", "BEST_GROWTH_POLICY", "BEST_SURVIVAL_POLICY"):
        pinfo.append({"item": k, "value": rv.get(k)})
    wf = rv.get("walk_forward", {}) if isinstance(rv.get("walk_forward"), dict) else {}
    pinfo += [{"item": "walk_forward_months_available", "value": wf.get("months_evaluated")},
              {"item": "walk_forward_months_required", "value": wf.get("required")},
              {"item": "walk_forward_status", "value": wf.get("status")}]
    T["Current_Policy"] = pd.concat([pd.DataFrame(pinfo).assign(section="policy"), pd.DataFrame(pr).assign(section="tracks"),
                                     pd.DataFrame(pb).assign(section="current-policy backtest view")], ignore_index=True)

    # meaningful cell leaderboards
    def lb(dfc, lo, hi, lab):
        x = dfc[(dfc["N"] >= lo) & (dfc["N"] <= hi)]
        parts = [x.nlargest(10, "roi").assign(rank_by="best ROI"), x.nsmallest(10, "roi").assign(rank_by="worst ROI"),
                 x.nlargest(10, "pnl_units").assign(rank_by="best PnL"), x.nsmallest(10, "pnl_units").assign(rank_by="worst PnL")]
        return pd.concat(parts, ignore_index=True).assign(sample_band=lab)
    T["Best_Worst_Cells"] = pd.concat([lb(cells, 50, 10 ** 9, "N>=50"), lb(cells, 25, 49, "N 25-49")], ignore_index=True)

    T["V11_League"], T["V11_Global"] = v11_tables(I, nonmls, end)
    T["Top_Picks"] = top_picks_table(I)
    T["Pure_Prediction"] = pure_table(I)
    T["Retrain_Health"] = retrain_table(I)
    T["Player_V3"] = pd.DataFrame([{"field": k, "value": json.dumps(v) if isinstance(v, (dict, list)) else v}
                                   for k, v in (I["player_v3"] or {}).items()]) if I["player_v3"] else _insufficient()
    T["League_Scout"] = scout_table(I)
    T.update(lineage_tables())
    try:
        T["Evidence_Matrix"] = evidence_matrix({**I, "v11_placebo": I.get("v11_placebo")})
    except Exception as e:                                             # noqa: BLE001
        T["Evidence_Matrix"] = _insufficient(f"evidence matrix failed: {type(e).__name__}: {e}")
    T["_meta"] = pd.DataFrame([{"window": win_txt}])
    return T


def evidence_matrix(I: dict) -> pd.DataFrame:
    """TOP-LEVEL EVIDENCE MATRIX (owner brief 2026-10-19 §19). Keeps six different things apart:
    live betting results, predictive skill, market-relative skill, price movement, CLV and executable
    economic edge. A cell can only become EXECUTABLE_EDGE with real economic validation — no row can
    reach it today, and the rules below make that structural rather than a matter of judgement."""
    s = I["s"]
    pool = s[~s["league"].isin(cfg.PAPER_LEAGUES) & s["tier"].isin(C.TIERS)]
    end = I["end"]
    resid = (_json(cfg.OUTPUT_DIR / "studies" / "ou_studies.json") or {}).get("residual", {})
    mk = {"OU25": "OU25", "OVER15": "OU15", "OVER35": "OU35", "BTTS": "BTTS"}
    plc = I.get("v11_placebo")
    move_txt = "n/a"
    if plc is not None and len(plc) and "variant" in plc.columns:
        try:
            v9r = float(plc.loc[plc["variant"] == "v9_residual", "toward_rate"].iloc[0])
            fx = float(plc.loc[plc["variant"] == "PLACEBO_fixed_anchor", "toward_rate"].iloc[0])
            move_txt = f"toward {v9r:.3f} vs fixed-anchor {fx:.3f} ({(v9r - fx) * 100:+.1f}pp)"
        except Exception:                                              # noqa: BLE001
            pass
    rows = []
    for (model, market), g in pool.groupby(["model", "market"]):
        if len(g) < 30:
            continue
        sm = C.summarize(g)
        lo, hi = sm.get("roi_ci_lo_90"), sm.get("roi_ci_hi_90")
        r = (resid.get(mk.get(market, market)) or {}).get(model, {})
        verdict = r.get("verdict") or r.get("status") or "NOT_TESTED"
        econ = ("INSUFFICIENT" if len(g) < 100 else "NEGATIVE" if (hi is not None and hi < 0) else
                "INTERESTING" if sm["roi"] > 0 else "NEGATIVE")
        mrel = {"MODEL_HURTS": "NEGATIVE", "NO_DETECTABLE_GAIN": "NONE_DETECTED",
                "TOO_FEW": "INSUFFICIENT"}.get(verdict, "MARKET_RELATIVE_SIGNAL" if "GAIN" in verdict else verdict)
        rows.append({"track": f"{model} | {market}", "N": len(g), "ROI": sm["roi"],
                     "ROI_CI90": f"[{lo:+.3f}, {hi:+.3f}]" if lo is not None else "n/a",
                     "CLV": sm.get("avg_clv"), "predictive_gain_vs_baseline":
                     "PREDICTIVE_SIGNAL (models beat base rate, AUC .57-.64 - Prediction Lab)",
                     "gain_vs_market": f"{mrel} ({verdict}, OOS dLL {r.get('oos_delta_logloss', 'n/a')})",
                     "movement_vs_placebo": move_txt if market == "OU25" else "not measured for this market",
                     "forward_test_status": "no registered forward test for this track",
                     "economic_status": econ,
                     "confidence": C.confidence_label(len(g))})
    pt = _json(cfg.OUTPUT_DIR / "price_trader" / "v2" / "verdict_v2.json") or {}
    hp = pt.get("historical_primary", {})
    rows.append({"track": "Price Trader (back-to-lay, paper)", "N": hp.get("n"), "ROI": hp.get("mean_pnl_per_unit"),
                 "ROI_CI90": "see output/price_trader/v2", "CLV": None, "predictive_gain_vs_baseline": "n/a",
                 "gain_vs_market": "n/a", "movement_vs_placebo": "Wowza side moves LESS often than every control",
                 "forward_test_status": pt.get("EVIDENCE_STATE", "n/a"),
                 "economic_status": "NEGATIVE" if (hp.get("mean_pnl_per_unit") or 0) < 0 else "INSUFFICIENT",
                 "confidence": "historical exploratory; forward from 2026-10-20"})
    v3 = _json(cfg.OUTPUT_DIR / "player_v3" / "forward_test.json") or {}
    rows.append({"track": "Player props V3 (paper only, permanently)", "N": v3.get("settled_props"),
                 "ROI": None, "ROI_CI90": "per market in FORWARD_TEST.md", "CLV": None,
                 "predictive_gain_vs_baseline": "PREDICTIVE_SIGNAL (AUC 0.62-0.85)",
                 "gain_vs_market": "MARKET_RELATIVE_SIGNAL in probability (Player V3 lab); not in money",
                 "movement_vs_placebo": "n/a", "forward_test_status": "COLLECTING (V1); V2 two-sided registered",
                 "economic_status": "INSUFFICIENT", "confidence": "forward bets < 150 in every market"})
    fb = _json(cfg.OUTPUT_DIR / "studies" / "fantasy_benchmark.json") or {}
    fpl = (fb.get("pooled_all_gws") or {})
    rows.append({"track": "Fantasy (vs FPL ep_next)", "N": fb.get("player_gameweeks"), "ROI": None,
                 "ROI_CI90": "n/a (points, not money)", "CLV": None,
                 "predictive_gain_vs_baseline": (f"MAE Wowza {fpl.get('WOWZA', {}).get('mae')} vs FPL "
                                                 f"{fpl.get('FPL_EP_NEXT', {}).get('mae')}"),
                 "gain_vs_market": f"adds beyond FPL: {fb.get('WOWZA_ADDS_INFORMATION_BEYOND_FPL', 'n/a')}",
                 "movement_vs_placebo": "n/a", "forward_test_status": "prospective challenger log running",
                 "economic_status": "NEGATIVE (FPL is the champion)", "confidence": "3 gameweeks"})
    ls = _json(cfg.OUTPUT_DIR / "longshot_plan" / "forward_status.json") or {}
    rows.append({"track": "Longshot BET_ENABLED filter v1", "N": None, "ROI": None,
                 "ROI_CI90": "see pre_freeze_months_EXPLORATORY.csv", "CLV": None,
                 "predictive_gain_vs_baseline": "n/a", "gain_vs_market": "n/a", "movement_vs_placebo": "n/a",
                 "forward_test_status": ls.get("STATUS", "n/a"),
                 "economic_status": "NEGATIVE (does not beat simple baselines)", "confidence": "frozen 2026-10-19"})
    out = pd.DataFrame(rows)
    out["EXECUTABLE_EDGE"] = "NO"          # no row has real executable economic validation
    return out


#: Champion high-water mark (owner 2026-10-18): sheet -> model_lineage output.
LINEAGE_SHEETS = {"Model_Lineage": "current_champions_view", "Retrain_Batches": "retrain_batch_results",
                  "Champion_History": "champion_live_performance",
                  "Challenger_Results": "challenger_shadow_results",
                  "Live_By_Model_Version": "model_version_live_results",
                  "Model_High_Water": "champion_health", "Model_Churn": "model_churn"}


def lineage_tables() -> dict[str, pd.DataFrame]:
    """The seven lineage sheets. Never blocks the workbook: a failure becomes INSUFFICIENT_DATA."""
    try:
        from src.analytics import model_lineage as ML
        r = ML.build(cfg.V9_LOCAL, results=ML.load_results())
        return {sheet: r[key] for sheet, key in LINEAGE_SHEETS.items()}
    except Exception as e:                                             # noqa: BLE001
        return {sheet: _insufficient(f"model_lineage failed: {type(e).__name__}: {e}")
                for sheet in LINEAGE_SHEETS}


def _insufficient(msg: str = "source missing") -> pd.DataFrame:
    return pd.DataFrame([{"STATUS": "INSUFFICIENT_DATA", "detail": msg}])


def v11_tables(I: dict, nonmls: pd.DataFrame, end) -> tuple[pd.DataFrame, pd.DataFrame]:
    L, R, P, S = I["v11_league"], I["v11_residual"], I["v11_placebo"], I["v11_summary"]
    if L.empty:
        return _insufficient("v11 outputs unreachable"), _insufficient("v11 outputs unreachable")
    bet = C.table(nonmls, ["league"], end, RECENT_FROM, ci=False)[["league", "N", "roi", "recent_ROI", "pnl_units"]]
    v = L.rename(columns={"segment": "league", "n_fixtures": "movement_fixtures", "mean_signed_move_pp": "mean_signed_move_pp",
                          "pct_positive_clv": "pct_positive_clv_v11"})
    v = v[["league", "movement_fixtures", "toward_rate", "toward_ci_lo", "toward_ci_hi", "mean_signed_move_pp",
           "mean_abs_move_pp", "mean_clv_pct", "median_clv_pct", "pct_positive_clv_v11", "sample_status"]]
    for c in ("mean_clv_pct", "median_clv_pct"):
        v[c] = pd.to_numeric(v[c], errors="coerce") / 100.0
    r = R.rename(columns={"scope": "league"})
    r["brier_delta"] = r["brier_market_wowza"] - r["brier_market"]
    r["logloss_delta"] = r["logloss_market_wowza"] - r["logloss_market"]
    r = r[["league", "n", "brier_market", "brier_market_wowza", "brier_delta", "logloss_market", "logloss_market_wowza",
           "logloss_delta", "wowza_helps"]].rename(columns={"n": "residual_n"})
    placebo = P[P["variant"] == "PLACEBO_fixed_anchor"]["toward_rate"]
    ptw = float(placebo.iloc[0]) if len(placebo) else np.nan
    j = bet.merge(v, on="league", how="outer").merge(r, on="league", how="left")
    j["placebo_fixed_anchor_toward"] = ptw
    j["toward_minus_placebo"] = j["toward_rate"] - ptw

    def classify(x):
        if (x.get("N") or 0) < 40 or pd.isna(x.get("toward_rate")):
            return "INSUFFICIENT_DATA"
        res = (x["roi"] or 0) > 0 and (x.get("recent_ROI") or 0) > 0
        price = (x["mean_clv_pct"] or 0) > 0 and (x["toward_minus_placebo"] or 0) > 0
        resid = bool(x.get("wowza_helps")) is True and (x.get("logloss_delta") or 0) < 0
        if res and price and resid:
            return "STRONG_CONVERGENCE"
        if not res and not price and not resid:
            return "NEGATIVE"
        if res and not price and not resid:
            return "RESULTS_ONLY"
        if price and not res and not resid:
            return "PRICE_ONLY"
        if resid and not res and not price:
            return "RESIDUAL_ONLY"
        return "MIXED"
    j["convergence_class"] = j.apply(classify, axis=1)
    j["note"] = "research-only classification; never a selection input"

    g = S[["segment", "unit", "n_obs", "n_fixtures", "toward_rate", "toward_ci_lo", "toward_ci_hi", "mean_signed_move_pp",
           "mean_abs_move_pp"] + [c for c in ("mean_clv_pct", "median_clv_pct", "pct_positive_clv") if c in S.columns]].copy()
    g = pd.concat([g, P.assign(segment=P["variant"])[["segment", "toward_rate", "n_obs", "n_fixtures"]].assign(unit="placebo/control"),
                   I["v11_time"][["segment", "unit", "n_fixtures", "toward_rate"]].assign(segment=lambda x: "time-to-KO " + x["segment"])],
                  ignore_index=True)
    h = I["v11_health"] or {}
    g = pd.concat([pd.DataFrame([{"segment": f"health: {k}", "unit": json.dumps(v)[:200] if isinstance(v, (dict, list)) else v}
                                 for k, v in h.items()]), g], ignore_index=True)
    return j, g


def top_picks_table(I: dict) -> pd.DataFrame:
    r = I["tip_ranker"]
    if not r:
        return _insufficient("tip_ranker.json missing")
    rows = [{"list": k, **v} for k, v in r.get("lists", {}).items()]
    cal = r.get("calibration", {})
    rows.append({"list": "CALIBRATION", "hit_rate": cal.get("realised"), "brier_ranker": cal.get("brier_ranker"),
                 "brier_bookmaker_implied": cal.get("brier_implied"), "brier_raw_model": cal.get("brier_model")})
    return pd.DataFrame(rows)


def pure_table(I: dict) -> pd.DataFrame:
    v = I["pure"]
    if not v:
        return _insufficient("pure_prediction/verdicts.json missing")
    rows = []
    for t, d in (v.get("_detail") or {}).items():
        if isinstance(d, dict):
            rows.append({"target": t, "N": d.get("n"), "accuracy": d.get("accuracy"), "baseline_accuracy": d.get("rock"),
                         "accuracy_lift_pp": d.get("lift_pp"), "logloss": d.get("log_loss"),
                         "baseline_logloss": d.get("rock_log_loss"), "brier": d.get("brier"), "auc": d.get("auc"),
                         "ece": d.get("ece"), "folds_won_on_logloss": d.get("folds_won_on_logloss")})
    keys = ("DATASET_LEAKAGE_SAFE", "MORE_DATA_IS_IMPROVING_MODELS", "MARKET_ONLY_STRONGER_THAN_FOOTBALL",
            "MARKET_ADDS_TO_FOOTBALL_MODEL", "FOOTBALL_ONLY_ADDS_SIGNAL", "RECENCY_WEIGHTING_HELPS", "LEAGUE_SPECIFIC_MODELS_HELP")
    rows += [{"target": f"VERDICT {k}", "accuracy": None, "value": v.get(k, "NOT_IN_SOURCE")} for k in keys]
    return pd.DataFrame(rows)


def retrain_table(I: dict) -> pd.DataFrame:
    j = I["retrain"]
    runs = (j or {}).get("runs") or {}
    key = j.get("latest_run") if j.get("latest_run") in runs else (list(runs)[-1] if runs else None)
    if not key:
        return _insufficient("retrain_log.json missing")
    rows = []
    for model, e in runs[key].items():
        g = e.get("v91_gate") or {}
        checks = g.get("checks") if isinstance(g.get("checks"), dict) else {}
        can = e.get("canary") or {}
        rows.append({"run": key, "model": model, "rows": e.get("rows"), "incumbent_logloss": e.get("logloss_old"),
                     "challenger_logloss": e.get("logloss_new"),
                     "delta": (e.get("logloss_new") - e.get("logloss_old")) if e.get("logloss_new") is not None and e.get("logloss_old") is not None else None,
                     "LIVE_GATE_PROMOTED": e.get("promoted"), "live_gate_why": e.get("why"),
                     "STRICT_GATE_PROVEN": g.get("would_promote", g.get("decision")),
                     "strict_gate_detail": json.dumps(checks)[:500] if checks else json.dumps(g)[:500],
                     "canary": json.dumps(can)[:300], "comparison_basis": e.get("comparison_basis"), "at": e.get("at")})
    return pd.DataFrame(rows)


def scout_table(I: dict) -> pd.DataFrame:
    r, st = I["scout"], I["scout_status"]
    if not r and not st:
        return _insufficient("scout outputs missing")
    rows = [{"item": "leagues_monitored", "value": r.get("leagues_monitored")}]
    rows += [{"item": k, "value": v} for k, v in (r.get("counts") or {}).items()]
    rows += [{"item": f"status {k}", "value": v} for k, v in (st.get("status_counts") or {}).items()]
    near = (r.get("details") or {})
    for name, d in list(near.items())[:40]:
        if isinstance(d, dict) and d.get("distance"):
            rows.append({"item": f"near: {name}", "value": json.dumps(d.get("distance"))[:300]})
    return pd.DataFrame(rows)


# ── validation ───────────────────────────────────────────────────────────────────────────────
def reconcile(I: dict) -> pd.DataFrame:
    """Recompute CORE a second, independent way (raw ledgers, groupby-first rather than first_sight())."""
    s, end, pol = I["s"], I["end"], I["policy"]
    core = C.in_window(C.universes(s, cfg.PAPER_LEAGUES, pol)["CORE"], C.windows(end)[f"since_{MAIN_START}"])
    raw = I["raw"]
    r = raw[raw["source"].eq("live") & raw["tier"].isin(C.CORE_TIERS) & ~raw["league"].isin(cfg.PAPER_LEAGUES)].copy()
    r = r.sort_values("generated_at", kind="mergesort")
    r = r.groupby(["day", "league", "home_team", "away_team", "market"], as_index=False).first()
    r = r[r["result"].isin(["WIN", "LOSS", "VOID", "PUSH"]) & (r["odds"] > 1) & (r["day"] >= MAIN_START) & (r["day"] <= str(end.date()))]
    alt_pnl = float(np.where(r["result"] == "WIN", r["odds"] - 1, np.where(r["result"] == "LOSS", -1.0, 0.0)).sum())
    checks = [("core N: first_sight vs groupby-first", len(core), len(r)),
              ("core PnL: first_sight vs groupby-first", round(core["pnl"].sum(), 6), round(alt_pnl, 6)),
              ("sum of model N = core N", int(core.groupby("model").size().sum()), len(core)),
              ("sum of model PnL = core PnL", round(core.groupby("model")["pnl"].sum().sum(), 6), round(core["pnl"].sum(), 6)),
              ("sum of market PnL = core PnL", round(core.groupby("market")["pnl"].sum().sum(), 6), round(core["pnl"].sum(), 6)),
              ("sum of tier PnL = core PnL", round(core.groupby("tier")["pnl"].sum().sum(), 6), round(core["pnl"].sum(), 6)),
              ("sum of league PnL = core PnL", round(core.groupby("league")["pnl"].sum().sum(), 6), round(core["pnl"].sum(), 6))]
    out = pd.DataFrame([{"check": c, "value_a": a, "value_b": b, "status": "PASS" if abs(float(a) - float(b)) < 1e-6 else "FAIL"}
                        for c, a, b in checks])
    led = core["ledger_pnl"]
    diff = (core["pnl"] - led).abs()
    out = pd.concat([out, pd.DataFrame([{"check": "info: flat PnL vs ledger pnl column (rows differing > 0.01)",
                                         "value_a": int((diff > 0.01).sum()), "value_b": int(led.notna().sum()),
                                         "status": "INFO"}])], ignore_index=True)
    return out


def data_quality(I: dict) -> pd.DataFrame:
    raw, st, fs = I["raw"], I["stats"], I["first"]
    live = raw[raw["source"].eq("live")]
    res = live["result"]
    rows = [{"item": k, "value": v} for k, v in st.items()]
    rows += [{"item": "settled rows (live, raw)", "value": int(res.isin(["WIN", "LOSS", "VOID", "PUSH"]).sum())},
             {"item": "pending rows (live, raw)", "value": int((res == "").sum())},
             {"item": "missing odds", "value": int(live["odds"].isna().sum())},
             {"item": "invalid odds (<= 1.0)", "value": int((live["odds"] <= 1.0).sum())},
             {"item": "missing closing odds (settled first-sight)", "value": int(I["s"]["closing_odds"].isna().sum())},
             {"item": "CLV coverage (settled first-sight)", "value": f"{I['s']['clv'].notna().mean():.1%}"},
             {"item": "missing model type", "value": int((live["model"] == "unknown").sum())},
             {"item": "side=AVOID rows (model declined; excluded, never a bet)", "value": int((live["side"] == "AVOID").sum())},
             {"item": "missing tier", "value": int(live["tier"].isna().sum())},
             {"item": "rows by source", "value": json.dumps(raw["source"].value_counts().to_dict())},
             {"item": "market names", "value": json.dumps(fs["market"].value_counts().to_dict())},
             {"item": "unparseable match dates", "value": int(pd.to_datetime(live["day"], errors="coerce").isna().sum())},
             {"item": "impossible ledger PnL (|pnl| > odds)", "value": int((live["ledger_pnl"].abs() > live["odds"]).sum())}]
    return pd.DataFrame(rows)


# ── workbook ─────────────────────────────────────────────────────────────────────────────────
SHEETS = ["Dashboard", "Evidence_Matrix", "Methodology", "Overall", "Market_Model", "Tier", "Model_Market_Tier", "League", "League_Market",
          "Detailed_Cells", "Odds_Bands", "Edge_Bands", "OU25_Side", "Monthly", "Weekly", "Daily", "Rolling",
          "Profit_Curves", "Drawdown", "V11_League", "V11_Global", "CLV", "Calibration", "Season_Projection",
          "Current_Policy", "Top_Picks", "Pure_Prediction", "Retrain_Health", "Player_V3", "League_Scout",
          "Model_Lineage", "Retrain_Batches", "Champion_History", "Challenger_Results",
          "Live_By_Model_Version", "Model_High_Water", "Model_Churn",
          "System_Health", "Data_Quality", "Definitions"]


def write_workbook(path: Path, T: dict, I: dict, A: dict, recon: pd.DataFrame, dq: pd.DataFrame) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    HEAD = PatternFill("solid", fgColor="1F2937")
    GREEN, RED, GREY = PatternFill("solid", fgColor="D1FAE5"), PatternFill("solid", fgColor="FEE2E2"), PatternFill("solid", fgColor="F3F4F6")
    win = f"Window: {MAIN_START} -> {I['end'].date()} (latest settled) unless the table says otherwise · first-sight bets · 1u flat stake"

    def put(name: str, df: pd.DataFrame, note: str = "", start_row: int = 3):
        ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        ws["A1"] = name.replace("_", " ")
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = note or win
        ws["A2"].font = Font(italic=True, color="6B7280")
        if df is None or df.empty:
            df = _insufficient()
        df = df.copy()
        for c in df.columns:
            if df[c].dtype == object:
                df[c] = df[c].map(lambda v: json.dumps(v) if isinstance(v, (list, dict)) else v)
        cols = list(df.columns)
        for j, c in enumerate(cols, start=1):
            cell = ws.cell(row=start_row, column=j, value=str(c))
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
            for j, v in enumerate(row, start=1):
                if isinstance(v, float) and (np.isnan(v) or np.isinf(v)):
                    v = None
                elif isinstance(v, (np.integer,)):
                    v = int(v)
                elif isinstance(v, (np.floating,)):
                    v = float(v)
                elif isinstance(v, pd.Timestamp):
                    v = v.to_pydatetime().replace(tzinfo=None)
                ws.cell(row=i, column=j, value=v)
        last = start_row + len(df)
        for j, c in enumerate(cols, start=1):
            L = get_column_letter(j)
            ws.column_dimensions[L].width = min(42, max(10, len(str(c)) + 2))
            fmt = "0.0%" if c in PCT_COLS or c.endswith("_ROI") or c.startswith("share") else (
                "0.00" if df[c].dtype.kind == "f" else None)
            if fmt:
                for i in range(start_row + 1, last + 1):
                    ws.cell(row=i, column=j).number_format = fmt
            if c in ("roi", "ROI", "recent_ROI", "current_ROI", "last30_ROI", "pnl_units", "PnL", "exARG_ROI", "exTOP_ROI"):
                rng = f"{L}{start_row + 1}:{L}{last}"
                ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=GREEN))
                ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=RED))
            if c == "confidence_label":
                rng = f"{L}{start_row + 1}:{L}{last}"
                for lab in ("VERY_SMALL", "SMALL"):
                    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{lab}"'], fill=GREY))
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{last}"
        return ws, start_row, last, cols

    # Dashboard
    ws = wb.create_sheet("Dashboard")
    ws["A1"], ws["A1"].font = "WOWZA — deep analytics", Font(bold=True, size=18)
    ws["A2"] = win
    r = 4
    for k, v in A["kpis"]:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=r, column=2, value=v)
        if isinstance(v, float) and abs(v) < 5 and ("ROI" in k or "CLV" in k or "toward" in k):
            c.number_format = "0.0%"
        r += 1
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 44
    r += 1
    for title, df in A["dash_tables"]:
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=12)
        r += 1
        for j, c in enumerate(df.columns, start=1):
            x = ws.cell(row=r, column=j, value=str(c)); x.font, x.fill = Font(bold=True, color="FFFFFF"), HEAD
        for row in df.itertuples(index=False):
            r += 1
            for j, v in enumerate(row, start=1):
                ws.cell(row=r, column=j, value=None if isinstance(v, float) and np.isnan(v) else v)
        r += 2

    put("Evidence_Matrix", T.get("Evidence_Matrix"),
        note=("Six different things, kept apart: live betting ROI / predictive skill / market-relative skill / "
              "price movement vs placebo / forward-test status / economic status. EXECUTABLE_EDGE requires real "
              "economic validation - no row has it."))
    put("Methodology", A["methodology"], note="How every number in this workbook is produced")
    for name in SHEETS[2:]:
        if name in ("System_Health", "Data_Quality", "Definitions", "Methodology"):
            continue
        df = T.get(name)
        note = DISCLAIMER if name == "Season_Projection" else ""
        put(name, df, note=note)
    put("Data_Quality", pd.concat([dq, recon.rename(columns={"check": "item"}).assign(value=lambda x: x["status"])], ignore_index=True))
    put("System_Health", A["system_health"])
    put("Definitions", A["definitions"], note="Terms used throughout")
    for extra in ("Edge_Rank", "Concentration", "Ex_Argentina_Top", "CLV_vs_ROI", "Daily_Volume", "Best_Worst_Cells"):
        put(extra, T.get(extra))
    put("Reconciliation", recon, note="Core recomputed a second, independent way + additivity checks")

    # charts
    def line_chart(sheet, title, df_name, ycol, at):
        df = T[df_name]
        ws = wb[sheet]
        hdr = 3
        for track, g in df.groupby("track"):
            pass
        ch = LineChart(); ch.title, ch.height, ch.width = title, 9, 20
        col = list(df.columns).index(ycol) + 1
        start = hdr + 1
        for track in df["track"].unique():
            idx = df.index[df["track"] == track]
            s = Series(Reference(ws, min_col=col, min_row=start + idx.min(), max_row=start + idx.max()), title=str(track))
            ch.series.append(s)
        ws.add_chart(ch, at)

    def bar_chart(sheet, title, cat_col, val_col, at, mask=None):
        df = T[sheet] if sheet in T else None
        ws = wb[sheet]
        if df is None or df.empty or val_col not in df.columns:
            return
        sub = df if mask is None else df[mask(df)]
        if sub.empty:
            return
        ci, vi = list(df.columns).index(cat_col) + 1, list(df.columns).index(val_col) + 1
        rows = [4 + i for i in sub.index]
        ch = BarChart(); ch.title, ch.height, ch.width = title, 8, 18
        ch.add_data(Reference(ws, min_col=vi, min_row=rows[0], max_row=rows[-1]), titles_from_data=False)
        ch.set_categories(Reference(ws, min_col=ci, min_row=rows[0], max_row=rows[-1]))
        ch.y_axis.number_format = "0%"
        ws.add_chart(ch, at)

    try:
        line_chart("Profit_Curves", "Cumulative flat-stake PnL by track", "Profit_Curves", "cumulative_pnl", "H4")
        bar_chart("Monthly", "Monthly ROI (overall core)", "month", "roi", "Z4", lambda d: d["track"] == "overall_core")
        bar_chart("Weekly", "Weekly ROI (core)", "week", "roi", "AE4")
        bar_chart("Odds_Bands", "ROI by odds band (all non-MLS)", "odds_band", "roi", "AD4", lambda d: d["view"] == "all non-MLS")
        bar_chart("Edge_Bands", "ROI by edge band (all non-MLS)", "edge_band", "roi", "AD4", lambda d: d["view"] == "all non-MLS")
        bar_chart("League", "League ROI (N >= 25)", "league", "roi", "AF4", lambda d: d["N"] >= 25)
        bar_chart("Market_Model", "ROI by model x market (core)", "market", "roi", "AH4", lambda d: d["universe"] == "CORE")
        rl = T["Rolling"]
        if "roi_30d" in rl.columns:
            line_chart("Rolling", "Rolling 30-day ROI", "Rolling", "roi_30d", "AD4")
        for sheet, xcol, ycol, title in (("V11_League", "toward_rate", "roi", "Toward-Wowza rate vs ROI"),
                                         ("CLV_vs_ROI", "avg_clv", "roi", "Cell CLV vs ROI (N>=25)")):
            df = T.get(sheet)
            if df is None or df.empty or xcol not in df.columns:
                continue
            ws = wb[sheet]
            sub = df[df[xcol].notna() & df[ycol].notna()]
            if len(sub) < 3:
                continue
            xi, yi = list(df.columns).index(xcol) + 1, list(df.columns).index(ycol) + 1
            ch = ScatterChart(); ch.title, ch.style, ch.height, ch.width = title, 13, 9, 16
            s = Series(Reference(ws, min_col=yi, min_row=4, max_row=3 + len(df)),
                       Reference(ws, min_col=xi, min_row=4, max_row=3 + len(df)), title="cells")
            s.marker.symbol, s.graphicalProperties.line.noFill = "circle", True
            ch.series.append(s)
            ws.add_chart(ch, f"{get_column_letter(len(df.columns) + 2)}4")
        v = T.get("V11_League")
        if v is not None and "mean_clv_pct" in v.columns:
            bar_chart("V11_League", "CLV by league (V11)", "league", "mean_clv_pct", "AF24")
    except Exception as e:                                           # noqa: BLE001 - charts never block the workbook
        print(f"[deep] chart skipped: {type(e).__name__}: {e}")

    order = [n for n in SHEETS if n in wb.sheetnames] + [n for n in wb.sheetnames if n not in SHEETS]
    wb._sheets = [wb[n] for n in order]
    wb.save(path)


# ── answers, report ──────────────────────────────────────────────────────────────────────────
def answers(T: dict, I: dict, recon: pd.DataFrame) -> dict:
    end = I["end"]
    o = T["Overall"]
    get = lambda u, w, k: o[(o.universe == u) & (o.window == w)][k].iloc[0]  # noqa: E731
    mm = T["Market_Model"]; core_mm = mm[mm.universe == "CORE"]
    big = core_mm[core_mm.N >= 100]
    lg = T["League"]; lgm = lg[lg.N >= 50]
    ob = T["Odds_Bands"]; oba = ob[(ob.view == "all non-MLS") & (ob.N >= 30)]
    tr = T["Tier"]; tt = tr[tr.view == "tier"]
    er = T["Edge_Rank"]
    sp = T["Season_Projection"].set_index("row")
    res = T["V11_League"]
    pl = I["v11_placebo"]
    v9r = pl[pl.variant == "v9_residual"]["toward_rate"]; fa = pl[pl.variant == "PLACEBO_fixed_anchor"]["toward_rate"]
    vres = I["v11_residual"]; ov = vres[vres["scope"] == "overall"]
    mono = "UNCLEAR"
    if len(er):
        m = er["spearman_edge_vs_return"].dropna()
        mono = "YES" if len(m) and (m > 0.05).all() else "NO" if len(m) and (m <= 0).any() else "UNCLEAR"
    return {
        "ANALYSIS_START": MAIN_START, "ANALYSIS_END": str(end.date()), "LATEST_SETTLED_DATE": str(end.date()),
        "RAW_ROWS": I["stats"]["RAW_ROWS"], "UNIQUE_FIRST_SIGHT_ROWS": I["stats"]["UNIQUE_FIRST_SIGHT_ROWS"],
        "DUPLICATES_REMOVED": I["stats"]["DUPLICATES_REMOVED"],
        "CONFLICTING_SIDE_ROWS_FOUND": I["stats"]["CONFLICTING_SIDE_ROWS_FOUND"],
        "CORE_N": int(get("CORE", f"since_{MAIN_START}", "N")), "CORE_PNL": round(float(get("CORE", f"since_{MAIN_START}", "pnl_units")), 2),
        "CORE_ROI": round(float(get("CORE", f"since_{MAIN_START}", "roi")), 4),
        "RECENT_CORE_ROI": round(float(get("CORE", "since_2026-09-24", "roi")), 4),
        "NEW_FORMAT_ROI": round(float(core_mm[core_mm.model == "new_format"]["pnl_units"].sum() / core_mm[core_mm.model == "new_format"]["N"].sum()), 4),
        "STANDARD_ROI": round(float(core_mm[core_mm.model == "standard"]["pnl_units"].sum() / core_mm[core_mm.model == "standard"]["N"].sum()), 4),
        "BEST_LARGE_TRACK": f"{big.loc[big.roi.idxmax(), 'model']}|{big.loc[big.roi.idxmax(), 'market']} ({big.roi.max():+.1%})" if len(big) else "n/a",
        "WORST_LARGE_TRACK": f"{big.loc[big.roi.idxmin(), 'model']}|{big.loc[big.roi.idxmin(), 'market']} ({big.roi.min():+.1%})" if len(big) else "n/a",
        "BEST_MEANINGFUL_LEAGUE": f"{lgm.loc[lgm.roi.idxmax(), 'league']} ({lgm.roi.max():+.1%}, N={int(lgm.loc[lgm.roi.idxmax(), 'N'])})" if len(lgm) else "n/a",
        "WORST_MEANINGFUL_LEAGUE": f"{lgm.loc[lgm.roi.idxmin(), 'league']} ({lgm.roi.min():+.1%}, N={int(lgm.loc[lgm.roi.idxmin(), 'N'])})" if len(lgm) else "n/a",
        "BEST_ODDS_BAND": f"{oba.loc[oba.roi.idxmax(), 'odds_band']} ({oba.roi.max():+.1%})" if len(oba) else "n/a",
        "WORST_ODDS_BAND": f"{oba.loc[oba.roi.idxmin(), 'odds_band']} ({oba.roi.min():+.1%})" if len(oba) else "n/a",
        "BEST_TIER": f"{tt.loc[tt.roi.idxmax(), 'tier']} ({tt.roi.max():+.1%})" if len(tt) else "n/a",
        "EDGE_IS_MONOTONIC_WITH_ROI": mono,
        "WOWZA_ADDS_TO_MARKET_V11": ("YES" if len(ov) and bool(ov["wowza_helps"].iloc[0]) else "NO") if len(ov) else "UNCLEAR",
        "GLOBAL_MOVEMENT_BEATS_PLACEBO": ("YES" if len(v9r) and len(fa) and v9r.iloc[0] > fa.iloc[0] else "NO") if len(v9r) and len(fa) else "UNCLEAR",
        "CURRENT_POLICY_STATUS": f"{I['policy'].get('policy_version')} ({I['policy'].get('status')})",
        "SEASON_EQ_CORE_TIPS_FULL_WINDOW": round(float(sp.loc["overall core", "full_window_projected_tips"]), 0),
        "SEASON_EQ_CORE_PNL_FULL_WINDOW": round(float(sp.loc["overall core", "full_window_projected_PnL"]), 1),
        "SEASON_EQ_CORE_TIPS_RECENT": round(float(sp.loc["overall core", "recent_projected_tips"]), 0),
        "SEASON_EQ_CORE_PNL_RECENT": round(float(sp.loc["overall core", "recent_projected_PnL"]), 1),
        "RECONCILIATION_STATUS": "PASS" if (recon[recon.status != "INFO"]["status"] == "PASS").all() else "FAIL",
        "PRODUCTION_SELECTION_CHANGED": "NO", "THRESHOLDS_CHANGED": "NO", "STAKING_CHANGED": "NO",
        "_convergence": res.groupby("convergence_class")["league"].apply(list).to_dict() if "convergence_class" in res else {},
    }


def methodology() -> pd.DataFrame:
    return pd.DataFrame([
        ("Sources", "v9 output/bets_ledger.csv (O/U 2.5), side_bets_ledger.csv (BTTS / O1.5 / O3.5), ht_ledger.csv (half-time); "
                    "v9 config/bankroll_policy.json; Pro output/ research artifacts; V11 public outputs."),
        ("First sight", "Earliest live tip by generated_at per (match_date, league, home, away) for O/U 2.5 and per "
                        "(match_date, league, home, away, market) for side / HT markets. A later opposite side is the model "
                        "flipping, dropped and counted (CONFLICTING_SIDE_ROWS_FOUND)."),
        ("Flat stake", "1 unit per bet: WIN odds-1, LOSS -1, VOID/PUSH 0. ROI = PnL / N. Hit = W / (W+L). Break-even = mean(1/odds)."),
        ("Universes", "ALL_LIVE (all tiers incl. MLS), NON_MLS, CORE (SNIPER+MARKSMAN, no paper leagues), VALUABLE, "
                      "CURRENT_ACTIVE_POLICY (core on tracks ACTIVE/ACTIVE_LOW in the current policy), PAPER_TRACKS, "
                      "BLOCKED_TRACKS, HT (no tier). Never pooled without a label."),
        ("Windows", f"Main {MAIN_START} -> latest settled. Recent regime from {RECENT_FROM}. Last-N-days windows end at "
                    "the latest settled date."),
        ("Uncertainty", f"Bootstrap over match days ({C.N_BOOT} resamples): bets on one day are not independent. "
                        "90% and 95% ROI intervals, 90% hit interval; only for N >= 10."),
        ("Confidence label", "VERY_SMALL <10, SMALL 10-24, EARLY 25-49, BUILDING 50-99, MEANINGFUL 100-199, MATURE 200+. Descriptive only."),
        ("Season projection", f"Season = {C.SEASON_DAYS} calendar days. tips = N / observed days x {C.SEASON_DAYS}; PnL = tips x ROI. "
                              f"Conservative: shrunk ROI = N/(N+{C.SHRINK_K}) x ROI. " + DISCLAIMER),
        ("V11 join", "League movement (toward-Wowza rate, signed move, CLV) and market residual (Brier / log loss with "
                     "and without Wowza). Toward rate is compared with the fixed-anchor PLACEBO, never with 50%."),
        ("Convergence", "STRONG_CONVERGENCE only when results (ROI and recent ROI > 0), price (V11 CLV > 0 and toward "
                        "above placebo) and residual (Wowza improves log loss) all agree and N >= 40. Research only."),
        ("Not changed", "No selection logic, threshold, model, staking, league eligibility or production behaviour."),
    ], columns=["topic", "method"])


def definitions() -> pd.DataFrame:
    return pd.DataFrame([
        ("N", "settled first-sight bets"), ("ROI", "PnL / N at 1 unit per bet"), ("PnL", "profit in units at 1 unit flat"),
        ("CLV", "closing-line value: entry price vs closing price, from the ledgers (positive = beat the close)"),
        ("toward Wowza", "share of fixtures where the market price moved toward Wowza's side (V11)"),
        ("signed move", "price move in probability points, signed so positive = toward Wowza"),
        ("market residual", "does adding Wowza to the de-vigged market improve Brier / log loss (V11)"),
        ("SNIPER / MARKSMAN / VALUABLE", "v9 signal tiers by edge; SNIPER+MARKSMAN are the staked tiers, VALUABLE is recorded"),
        ("ACTIVE / ACTIVE_LOW / PAPER / BLOCKED", "bankroll-policy track states: staked / staked small / recorded only / not even paper"),
        ("first sight", "the earliest tip for a fixture (and market); later updates are not new bets"),
        ("season equivalent", f"N per observed day x {C.SEASON_DAYS} days"), ("recent regime", f"from {RECENT_FROM}"),
        ("confidence label", "sample-size band, descriptive only"), ("shrunk ROI", f"N/(N+{C.SHRINK_K}) x observed ROI"),
    ], columns=["term", "definition"])


def report_md(A: dict, T: dict) -> str:
    a = A["answers"]
    mm = T["Market_Model"]; c = mm[mm.universe == "CORE"].sort_values("pnl_units", ascending=False)
    L = [f"# WOWZA deep analytics — {a['ANALYSIS_END']}", "",
         f"Window {a['ANALYSIS_START']} -> {a['ANALYSIS_END']} · first-sight bets · 1 unit flat. Research only.", "",
         "## Current state", f"Core: {a['CORE_N']} bets, {a['CORE_PNL']:+.2f}u, ROI {a['CORE_ROI']:+.2%}; since "
         f"{RECENT_FROM}: {a['RECENT_CORE_ROI']:+.2%}. New format {a['NEW_FORMAT_ROI']:+.2%}, standard {a['STANDARD_ROI']:+.2%}.", "",
         "## What is making money", *[f"- {r.model} {r.market}: {int(r.N)} bets, {r.pnl_units:+.2f}u ({r.roi:+.1%}), {r.confidence_label}"
                                      for r in c.head(4).itertuples() if r.pnl_units > 0], "",
         "## What is losing money", *[f"- {r.model} {r.market}: {int(r.N)} bets, {r.pnl_units:+.2f}u ({r.roi:+.1%}), {r.confidence_label}"
                                      for r in c.tail(4).itertuples() if r.pnl_units < 0], "",
         "## Odds, edge, tiers", f"- Best odds band: {a['BEST_ODDS_BAND']}; worst: {a['WORST_ODDS_BAND']}.",
         f"- Edge monotonic with ROI: {a['EDGE_IS_MONOTONIC_WITH_ROI']}. Best tier: {a['BEST_TIER']}.", "",
         "## Leagues and V11", f"- Best meaningful league: {a['BEST_MEANINGFUL_LEAGUE']}; worst: {a['WORST_MEANINGFUL_LEAGUE']}.",
         f"- V11: Wowza adds to market {a['WOWZA_ADDS_TO_MARKET_V11']}; movement beats placebo {a['GLOBAL_MOVEMENT_BEATS_PLACEBO']}.",
         *[f"- {k}: {', '.join(v)}" for k, v in a["_convergence"].items()], "",
         "## Season run-rate (scenarios, not forecasts)",
         f"- Full window: ~{a['SEASON_EQ_CORE_TIPS_FULL_WINDOW']:.0f} core tips, {a['SEASON_EQ_CORE_PNL_FULL_WINDOW']:+.1f}u.",
         f"- Recent regime: ~{a['SEASON_EQ_CORE_TIPS_RECENT']:.0f} tips, {a['SEASON_EQ_CORE_PNL_RECENT']:+.1f}u.", f"- {DISCLAIMER}", "",
         "## Final answer", "```", *[f"{k} = {v}" for k, v in a.items() if not k.startswith("_")], "```"]
    return "\n".join(L) + "\n"


def assemble(I: dict, T: dict) -> dict:
    recon = reconcile(I)
    dq = data_quality(I)
    a = answers(T, I, recon)
    o = T["Overall"]; ow = o[o.window == f"since_{MAIN_START}"].set_index("universe")
    mm = T["Market_Model"]; c = mm[mm.universe == "CORE"]
    roi_of = lambda mask: float(c[mask]["pnl_units"].sum() / c[mask]["N"].sum()) if c[mask]["N"].sum() else np.nan  # noqa: E731
    tt = T["Tier"][T["Tier"].view == "tier"].set_index("tier")
    pl = I["v11_placebo"]; v9r = pl[pl.variant == "v9_residual"]
    vs = I["v11_summary"]; vo = vs[vs.segment.astype(str).str.startswith("overall (fixture")]
    kpis = [("Analysis cutoff", a["ANALYSIS_END"]), ("Core bets", a["CORE_N"]), ("Core PnL (u)", a["CORE_PNL"]),
            ("Core ROI", a["CORE_ROI"]), ("Recent ROI (since 2026-09-24)", a["RECENT_CORE_ROI"]),
            ("New-format ROI", a["NEW_FORMAT_ROI"]), ("Standard ROI", a["STANDARD_ROI"]),
            ("BTTS ROI", roi_of(c.market == "BTTS")), ("OU25 ROI", roi_of(c.market == "OU25")),
            ("Sniper ROI", float(tt.loc["SNIPER", "roi"]) if "SNIPER" in tt.index else None),
            ("Marksman ROI", float(tt.loc["MARKSMAN", "roi"]) if "MARKSMAN" in tt.index else None),
            ("Valuable ROI", float(tt.loc["VALUABLE", "roi"]) if "VALUABLE" in tt.index else None),
            ("V11 toward (fixture level)", float(vo["toward_rate"].iloc[0]) if len(vo) else None),
            ("V11 toward, v9 residual (placebo table)", float(v9r["toward_rate"].iloc[0]) if len(v9r) else None),
            ("V11 residual state (Wowza adds to market)", a["WOWZA_ADDS_TO_MARKET_V11"]),
            ("Current policy", a["CURRENT_POLICY_STATUS"]), ("Reconciliation", a["RECONCILIATION_STATUS"])]
    cells = T["Detailed_Cells"]; m50 = cells[cells.N >= 50]
    cols = ["league", "model", "market", "tier", "N", "roi", "pnl_units", "confidence_label"]
    lg = T["League"]; lgm = lg[lg.N >= 50]
    dash = [("Top 5 positive meaningful cells (N>=50)", m50.nlargest(5, "roi")[cols]),
            ("Top 5 negative meaningful cells (N>=50)", m50.nsmallest(5, "roi")[cols]),
            ("Top 5 leagues (N>=50)", lgm.nlargest(5, "roi")[["league", "N", "roi", "pnl_units"]]),
            ("Most dangerous league (N>=50, lowest ROI)", lgm.nsmallest(1, "roi")[["league", "N", "roi", "pnl_units"]]),
            ("Most dangerous track (core, N>=50)", c[c.N >= 50].nsmallest(1, "roi")[["model", "market", "N", "roi", "pnl_units"]])]
    sh = [("v9 latest settled date", a["LATEST_SETTLED_DATE"]), ("v9 retrain latest run", (I["retrain"] or {}).get("latest_run")),
          ("bankroll review", (I["bankroll_rec"] or {}).get("generated_at")), ("Player V3 verdict", (I["player_v3"] or {}).get("generated_at")),
          ("scout review", (I["scout"] or {}).get("generated_at")), ("tip ranker", (I["tip_ranker"] or {}).get("generated_at")),
          ("V11 health", json.dumps((I["v11_health"] or {}).get("generated_at"))), ("V11 movement rows", len(I["v11_league"]))]
    return {"answers": a, "kpis": kpis, "dash_tables": dash, "recon": recon, "dq": dq, "methodology": methodology(),
            "definitions": definitions(), "system_health": pd.DataFrame(sh, columns=["component", "freshness / state"])}


def run(out_dir: Path | None = None) -> dict:
    I = load_inputs()
    T = build_tables(I)
    A = assemble(I, T)
    day = str(I["end"].date())
    out = Path(out_dir or cfg.OUTPUT_DIR / "deep_analytics" / day)
    out.mkdir(parents=True, exist_ok=True)
    write_workbook(out / f"WOWZA_Deep_Analytics_{day}.xlsx", T, I, A, A["recon"], A["dq"])
    (out / f"WOWZA_Deep_Analytics_{day}.md").write_text(report_md(A, T), encoding="utf-8")
    for name, f in (("Overall", "overall"), ("Market_Model", "market_model"), ("Tier", "tier"), ("League", "league"),
                    ("Detailed_Cells", "detailed_cells"), ("Odds_Bands", "odds_bands"), ("Edge_Bands", "edge_bands"),
                    ("V11_League", "v11_league"), ("Season_Projection", "season_projection")):
        T[name].to_csv(out / f"{f}.csv", index=False)
    (out / "answers.json").write_text(json.dumps({k: v for k, v in A["answers"].items()}, indent=1, default=str), encoding="utf-8")
    return {"out": out, **A}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    r = run(Path(a.out) if a.out else None)
    print(f"[deep] wrote {r['out']}")
    for k, v in r["answers"].items():
        if not k.startswith("_"):
            print(f"{k} = {v}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
